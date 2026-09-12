#!/usr/bin/env python3
"""親と子のつながりを説明する PowerPoint を生成する。"""
from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1B, 0x3A, 0x4B)
TEAL = RGBColor(0x0D, 0x73, 0x77)
ORANGE = RGBColor(0xC4, 0x6B, 0x2E)
RED = RGBColor(0xA6, 0x3D, 0x40)
INK = RGBColor(0x2C, 0x33, 0x38)
MUTED = RGBColor(0x5C, 0x6B, 0x73)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF6, 0xF1, 0xE8)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD5, 0xCC, 0xBE)
EA_FONT = "Yu Gothic"
LATIN_FONT = "Calibri"


def _set_ea_font(run, size_pt: float, bold: bool, color: RGBColor, font=EA_FONT):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = LATIN_FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", font)
    latin = rPr.find(qn("a:latin"))
    if latin is None:
        latin = etree.SubElement(rPr, qn("a:latin"))
    latin.set("typeface", LATIN_FONT)


def _box(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.adjustments[0] = 0.08
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    return sh


def _rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def _text(slide, l, t, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_ea_font(run, size, bold, color)
    return tb


def _bullets(slide, l, t, w, h, items, size=16, color=INK):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = item
        _set_ea_font(run, size, False, color)
    return tb


def _footer(slide, n, total=10):
    _text(
        slide, Inches(0.5), Inches(7.05), Inches(10), Inches(0.3),
        f"presence-logger  ・  親と子はどうつながるか  ・  {n} / {total}",
        size=11, color=MUTED,
    )


def _header_bar(slide, title, kicker=""):
    _rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.05), NAVY)
    _rect(slide, Inches(0), Inches(1.05), Inches(13.333), Inches(0.08), TEAL)
    if kicker:
        _text(slide, Inches(0.5), Inches(0.08), Inches(12), Inches(0.28), kicker, 12, False, RGBColor(0xA8, 0xD5, 0xD7))
        _text(slide, Inches(0.5), Inches(0.32), Inches(12), Inches(0.6), title, 28, True, WHITE)
    else:
        _text(slide, Inches(0.5), Inches(0.28), Inches(12), Inches(0.6), title, 28, True, WHITE)


def _device(slide, l, t, w, h, title, lines, fill=NAVY):
    _box(slide, l, t, w, h, fill)
    _text(slide, l + Inches(0.12), t + Inches(0.08), w - Inches(0.2), Inches(0.35), title, 14, True, WHITE)
    _bullets(
        slide, l + Inches(0.12), t + Inches(0.42), w - Inches(0.2), h - Inches(0.5),
        lines, size=12, color=WHITE,
    )


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def build(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 表紙
    s = new_slide(prs)
    _rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), NAVY)
    _rect(s, Inches(0), Inches(0), Inches(0.22), Inches(7.5), TEAL)
    _text(s, Inches(0.8), Inches(1.7), Inches(11), Inches(0.4), "presence-logger", 16, False, RGBColor(0xA8, 0xD5, 0xD7))
    _text(s, Inches(0.8), Inches(2.15), Inches(11.5), Inches(1.4), "親と子はどうつながっているか", 40, True, WHITE)
    _text(
        s, Inches(0.8), Inches(3.7), Inches(11), Inches(1.2),
        "子の IP は DHCP で変わります。つながりの本体は Wi-Fi と MAC と名前です。\nこの資料の数字は、ハブ初期設定の通し（ハブ tpc12345）の具体例です。",
        18, False, RGBColor(0xD7, 0xE6, 0xE7),
    )
    _text(s, Inches(0.8), Inches(6.6), Inches(11), Inches(0.3), "現場説明用  ・  IP を名簿に書かない理由", 14, False, RGBColor(0xA8, 0xD5, 0xD7))

    # 2 結論
    s = new_slide(prs)
    _header_bar(s, "結論：IP ではなく、3つでつながる")
    cards = [
        ("1. 電波", TEAL, "子はハブの Wi-Fi に入る", "SSID tpc12345-hub\nパスワードは Wi-Fi 用\nハブはいつも 10.42.0.1"),
        ("2. 本人確認", ORANGE, "機械の指紋は MAC", "再起動しても変わらない\nクローンでも複製されない\nIP が変わっても同一機"),
        ("3. 呼び方", NAVY, "名簿は名前だけ", "zero2 / tpc12345-002.local\nIP は書かない\n使う直前に今の番号を引く"),
    ]
    for i, (h, col, sub, body) in enumerate(cards):
        l = Inches(0.5 + i * 4.2)
        _box(s, l, Inches(1.45), Inches(3.95), Inches(5.15), CARD, LINE)
        _rect(s, l, Inches(1.45), Inches(3.95), Inches(0.12), col)
        _text(s, l + Inches(0.25), Inches(1.75), Inches(3.45), Inches(0.4), h, 20, True, col)
        _text(s, l + Inches(0.25), Inches(2.2), Inches(3.45), Inches(0.7), sub, 16, True, INK)
        _bullets(s, l + Inches(0.25), Inches(3.0), Inches(3.45), Inches(3.2), body.split("\n"), 15, MUTED)
    _footer(s, 2)

    # 3 登場人物
    s = new_slide(prs)
    _header_bar(s, "この通しの登場人物", "ハブ初期設定の机上確認で使った値")
    _device(
        s, Inches(0.5), Inches(1.4), Inches(4.1), Inches(3.3), "新ハブ tpc12345",
        ["工場IP 172.22.13.21", "子用 Wi-Fi tpc12345-hub", "ハブ自身 10.42.0.1 固定", "SSH は公開鍵"],
    )
    _device(
        s, Inches(4.85), Inches(1.4), Inches(4.1), Inches(3.3), "子A  zero2（引っ越し）",
        ["経路 1-1 で y", "名前・局番号そのまま", "MAC 2c:cf:67:aa:aa:aa", "朝の IP 10.42.0.47"],
        TEAL,
    )
    _device(
        s, Inches(9.2), Inches(1.4), Inches(3.65), Inches(3.3), "子B  tpc12345-002（新規）",
        ["経路 2 のクローン", "旧名は pizero2w-2", "MAC 88:a2:9e:bb:bb:bb", "朝の IP 10.42.0.80"],
        ORANGE,
    )
    _box(s, Inches(0.5), Inches(4.95), Inches(12.35), Inches(1.7), CREAM, LINE)
    _text(s, Inches(0.75), Inches(5.1), Inches(12), Inches(0.35), "動かないものと、変わってよいもの", 16, True, NAVY)
    _text(
        s, Inches(0.75), Inches(5.5), Inches(12), Inches(0.95),
        "動かない　ハブ名 / Wi-Fi名 / ハブの 10.42.0.1 / 各子の MAC / 付けたあとの名前\n"
        "変わる　　子の 10.42.0.x（DHCP）。朝 47 番でも、再起動後は 91 番でもよい。",
        16, False, INK,
    )
    _footer(s, 3)

    # 4 日常 子→親
    s = new_slide(prs)
    _header_bar(s, "日常：子から親へ記録を送る", "名前は使わない。送り先はハブの固定アドレス")
    _device(s, Inches(0.5), Inches(1.5), Inches(3.6), Inches(2.5), "子A zero2", ["MAC …:aa:aa:aa", "今の IP 10.42.0.47"], TEAL)
    _device(s, Inches(0.5), Inches(4.2), Inches(3.6), Inches(2.4), "子B tpc12345-002", ["MAC …:bb:bb:bb", "今の IP 10.42.0.80"], ORANGE)
    arr = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.3), Inches(3.15), Inches(3.1), Inches(0.7))
    arr.fill.solid()
    arr.fill.fore_color.rgb = TEAL
    arr.line.fill.background()
    _text(s, Inches(4.35), Inches(2.55), Inches(3.1), Inches(0.5), "MQTT 記録", 14, True, TEAL, PP_ALIGN.CENTER)
    _text(s, Inches(4.35), Inches(3.9), Inches(3.1), Inches(0.7), "Wi-Fi tpc12345-hub", 14, False, MUTED, PP_ALIGN.CENTER)
    _device(s, Inches(7.7), Inches(2.35), Inches(5.1), Inches(3.0), "ハブ tpc12345", ["受け口  10.42.0.1 :1883", "子の自分の番号は使わない", "ハブ IP は変わらない"])
    _footer(s, 4)

    # 5 親→子 SSH
    s = new_slide(prs)
    _header_bar(s, "親から子を操作するとき", "名簿に IP は無い。名前から「今の番号」を引く")
    steps = [
        ("1", "名簿を見る", "fleet/children.conf\nzero2\ntpc12345-002.local"),
        ("2", "今の IP を引く", "mDNS（.local）\nまたは AP の隣近所一覧"),
        ("3", "SSH する", "ssh pi@10.42.0.47\n公開鍵。パスワードなし"),
        ("4", "本人か確認", "MAC が 2c:cf:67:aa:aa:aa\nなら子A で確定"),
    ]
    for i, (n, title, body) in enumerate(steps):
        l = Inches(0.45 + i * 3.2)
        _box(s, l, Inches(1.5), Inches(3.0), Inches(4.55), CARD, LINE)
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, l + Inches(0.2), Inches(1.7), Inches(0.45), Inches(0.45))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        _text(s, l + Inches(0.2), Inches(1.74), Inches(0.45), Inches(0.4), n, 16, True, WHITE, PP_ALIGN.CENTER)
        _text(s, l + Inches(0.2), Inches(2.3), Inches(2.6), Inches(0.5), title, 18, True, NAVY)
        _bullets(s, l + Inches(0.2), Inches(2.95), Inches(2.6), Inches(2.7), body.split("\n"), 15, MUTED)
    _footer(s, 5)

    # 6 再起動
    s = new_slide(prs)
    _header_bar(s, "昼：子A を再起動した", "IP が変わっても、本人（MAC）と呼び方は同じ")
    _box(s, Inches(0.5), Inches(1.45), Inches(6.0), Inches(5.15), CARD, LINE)
    _text(s, Inches(0.75), Inches(1.6), Inches(5.5), Inches(0.4), "朝", 18, True, TEAL)
    _bullets(
        s, Inches(0.75), Inches(2.15), Inches(5.5), Inches(4.1),
        [
            "Wi-Fi　tpc12345-hub",
            "名前　zero2",
            "MAC　2c:cf:67:aa:aa:aa",
            "IP　10.42.0.47",
            "記録の送り先　10.42.0.1",
        ],
        18,
    )
    _box(s, Inches(6.85), Inches(1.45), Inches(6.0), Inches(5.15), CARD, LINE)
    _text(s, Inches(7.1), Inches(1.6), Inches(5.5), Inches(0.4), "再起動のあと", 18, True, ORANGE)
    _bullets(
        s, Inches(7.1), Inches(2.15), Inches(5.5), Inches(4.1),
        [
            "Wi-Fi　同じ tpc12345-hub",
            "名前　同じ zero2",
            "MAC　同じ 2c:cf:67:aa:aa:aa",
            "IP　10.42.0.91　← ここだけ変わる",
            "送り先　同じ 10.42.0.1",
        ],
        18,
    )
    _footer(s, 6)

    # 7 名前の前
    s = new_slide(prs)
    _header_bar(s, "名前ができる前（クローン起動直後）", "名簿の .local はまだ使わない。今この AP に居る MAC と IP")
    story = [
        ("1 電波だけ先", "SD に書いた tpc12345-hub でハブへ入る。仮の IP は 10.42.0.80。送り先はもう 10.42.0.1。"),
        ("2 隣近所を見る", "ハブの子用アンテナに「IP 10.42.0.80 / MAC 88:a2:9e:bb:bb:bb」と出る。名簿の zero2 とは MAC が違うので未登録。"),
        ("3 その IP に鍵で入る", "ssh pi@10.42.0.80（公開鍵）。hostname と聞くと古い名 pizero2w-2。これは名札の控えで、名簿のキーではない。"),
        ("4 同名でも迷子にならない", "旧親には元の pizero2w-2 が残っている。クローンは名前では探さない。今この Wi-Fi に居る MAC だけが本体。"),
    ]
    for i, (title, body) in enumerate(story):
        top = Inches(1.35 + i * 1.35)
        _box(s, Inches(0.5), top, Inches(12.35), Inches(1.22), CARD, LINE)
        _text(s, Inches(0.75), top + Inches(0.12), Inches(12), Inches(0.35), title, 16, True, TEAL)
        _text(s, Inches(0.75), top + Inches(0.48), Inches(12), Inches(0.6), body, 15, False, INK)
    _footer(s, 7)

    # 8 登録
    s = new_slide(prs)
    _header_bar(s, "名前が付くまで（経路 2 の登録）", "改名のあと IP はまた変わる。待つのは MAC")
    cols = [
        ("いま", "旧名 pizero2w-2\nIP 10.42.0.80\nMAC …:bb:bb:bb\n名簿には未掲載"),
        ("作業", "送信を止める\n局番号を空にする\n名前を tpc12345-002 に\n再起動"),
        ("復帰", "IP は 10.42.0.12 など別番号\nハブは MAC …:bb:bb:bb を待つ\n戻ってから名簿へ書く"),
        ("以降", "名簿 tpc12345-002.local\n今の IP は都度引き直す\n局番号は子の画面で付ける"),
    ]
    for i, (h, body) in enumerate(cols):
        l = Inches(0.45 + i * 3.2)
        _box(s, l, Inches(1.5), Inches(3.0), Inches(5.1), CARD, LINE)
        _rect(s, l, Inches(1.5), Inches(3.0), Inches(0.7), TEAL if i < 3 else NAVY)
        _text(s, l, Inches(1.58), Inches(3.0), Inches(0.55), h, 18, True, WHITE, PP_ALIGN.CENTER)
        _bullets(s, l + Inches(0.2), Inches(2.45), Inches(2.6), Inches(3.8), body.split("\n"), 15, INK)
    _footer(s, 8)

    # 9 1-1 vs 2
    s = new_slide(prs)
    _header_bar(s, "引っ越し（1-1）と新規（2）は、名簿の扱いが違う")
    _box(s, Inches(0.5), Inches(1.45), Inches(6.05), Inches(5.15), CARD, LINE)
    _rect(s, Inches(0.5), Inches(1.45), Inches(6.05), Inches(0.7), TEAL)
    _text(s, Inches(0.5), Inches(1.55), Inches(6.05), Inches(0.5), "1-1  動いている子を移す", 20, True, WHITE, PP_ALIGN.CENTER)
    _bullets(
        s, Inches(0.75), Inches(2.4), Inches(5.55), Inches(3.9),
        [
            "例：zero2",
            "旧親 raspberrypi5 の名簿に、すでに名前がある",
            "工場網で旧親へ SSH → 子の MAC を聞く",
            "Wi-Fi を tpc12345-hub へ切り替え",
            "新ハブの名簿へ追加し、旧親の名簿から消す",
            "名前も局番号もそのまま",
        ],
        16,
    )
    _box(s, Inches(6.8), Inches(1.45), Inches(6.05), Inches(5.15), CARD, LINE)
    _rect(s, Inches(6.8), Inches(1.45), Inches(6.05), Inches(0.7), ORANGE)
    _text(s, Inches(6.8), Inches(1.55), Inches(6.05), Inches(0.5), "2  クローンして増やす", 20, True, WHITE, PP_ALIGN.CENTER)
    _bullets(
        s, Inches(7.05), Inches(2.4), Inches(5.55), Inches(3.9),
        [
            "例：tpc12345-002",
            "コピーした SD から増やした別の機械",
            "旧親の名簿は触らない",
            "元の pizero2w-2 は旧親に残る（1-1 で n したため）",
            "新ハブの名簿にだけ追加",
            "改名し、局番号は空",
        ],
        16,
    )
    _footer(s, 9)

    # 10 まとめ表
    s = new_slide(prs)
    _header_bar(s, "まとめ：何でつながっているか")
    rows = [
        ("役割", "何でつなぐ", "この通しの例", True),
        ("電波・記録の送り先", "ハブ AP の SSID とパスワード。親は 10.42.0.1", "tpc12345-hub → 10.42.0.1:1883", False),
        ("どの機械か（本人）", "MAC。再起動でもクローンでも変わらない", "子A 2c:cf:67:aa:aa:aa", False),
        ("名簿・SSH の宛先", "ホスト名 / .local。IP は書かない", "zero2  /  tpc12345-002.local", False),
        ("名前ができる前", "今この AP に居る MAC と、そのときの IP", "10.42.0.80 と MAC …:bb:bb:bb", False),
        ("SSH の鍵", "公開鍵。Wi-Fi パスワードではない", "~/.ssh/id_ed25519", False),
    ]
    top = Inches(1.35)
    for i, (a, b, c, head) in enumerate(rows):
        y = top + Inches(i * 0.85)
        bg = NAVY if head else (CREAM if i % 2 == 0 else WHITE)
        fg = WHITE if head else INK
        _rect(s, Inches(0.45), y, Inches(2.6), Inches(0.8), bg)
        _rect(s, Inches(3.05), y, Inches(5.4), Inches(0.8), bg)
        _rect(s, Inches(8.45), y, Inches(4.4), Inches(0.8), bg)
        _text(s, Inches(0.55), y + Inches(0.2), Inches(2.4), Inches(0.5), a, 13, True, fg)
        _text(s, Inches(3.15), y + Inches(0.12), Inches(5.2), Inches(0.6), b, 13, False, fg)
        _text(s, Inches(8.55), y + Inches(0.2), Inches(4.2), Inches(0.5), c, 13, False, fg)
    _footer(s, 10)

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))


if __name__ == "__main__":
    out = Path(__file__).resolve().parents[1] / "docs" / "hub-and-child-connection.pptx"
    build(out)
    print(out)
